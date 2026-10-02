"""
Unit tests for the complete Document Ingestion Pipeline.
Tests extraction, cleaning, chunking, embedding, storage, and transaction safety.
External R2 and Gemini calls are mocked.
"""

import io
import uuid
from unittest.mock import MagicMock, patch
import docx
import pytest
from sqlalchemy import select

from backend.app.models.document import Document
from backend.app.models.chunk import DocumentChunk
from backend.app.models.user import User
from backend.app.services.ingestion import process_document_pipeline
from backend.app.rag.extractors import get_extractor
from backend.app.core.errors import AppException


def test_docx_extractor_with_headings():
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("Trace is an AI-powered technical documentation assistant.")
    doc.add_heading("Storage Details", level=2)
    doc.add_paragraph("Original documents are stored in Cloudflare R2.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = get_extractor("docx")
    units = extractor.extract(docx_bytes, "test.docx")

    assert len(units) >= 2
    assert units[0].section_title == "Architecture Overview"
    assert "technical documentation" in units[0].content
    assert units[1].section_title == "Storage Details"
    assert "Cloudflare R2" in units[1].content


def test_pipeline_success(test_db):
    user_id = uuid.uuid4()
    user = User(id=user_id, email=f"user_{user_id.hex[:6]}@example.com", display_name="Test User")
    test_db.add(user)
    test_db.commit()

    doc_id = uuid.uuid4()
    doc = Document(
        id=doc_id,
        owner_id=user_id,
        filename="guide.md",
        file_type="md",
        file_size=200,
        content_hash="dummy_hash_1",
        r2_object_key=f"documents/{doc_id}/guide.md",
        status="processing",
    )
    test_db.add(doc)
    test_db.commit()

    content = b"# Complete Guide\n\nThis is a full guide to the documentation system.\n\n## Section Two\nHere are additional details regarding data flow and operations."

    with patch("backend.app.services.ingestion.SessionLocal", return_value=test_db), \
         patch("backend.app.services.gemini.GeminiService.embed_documents", return_value=[[0.1] * 768] * 2):
        
        process_document_pipeline(
            document_id=doc_id,
            file_bytes=content,
            filename="guide.md",
            file_type="md",
        )

        doc = test_db.scalar(select(Document).where(Document.id == doc_id))
        assert doc.status == "ready"
        assert doc.chunk_count > 0
        assert doc.processing_error is None

        chunks = test_db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == doc_id)).all()
        assert len(chunks) == doc.chunk_count
        assert len(chunks[0].embedding) == 768


def test_pipeline_r2_download_failure_marks_failed(test_db):
    doc_id = uuid.uuid4()
    doc = Document(
        id=doc_id,
        owner_id=uuid.uuid4(),
        filename="missing.pdf",
        file_type="pdf",
        file_size=100,
        content_hash="dummy_hash_missing",
        r2_object_key=f"documents/{doc_id}/missing.pdf",
        status="processing",
    )
    test_db.add(doc)
    test_db.commit()

    mock_storage = MagicMock()
    mock_storage.download_file.side_effect = AppException(
        status_code=404, code="OBJECT_NOT_FOUND", message="Object missing"
    )

    with patch("backend.app.services.ingestion.SessionLocal", return_value=test_db), \
         patch("backend.app.services.ingestion.get_storage_service", return_value=mock_storage):
        
        process_document_pipeline(document_id=doc_id, file_bytes=None)

        doc = test_db.scalar(select(Document).where(Document.id == doc_id))
        assert doc.status == "failed"
        assert "missing" in doc.processing_error.lower() or "not found" in doc.processing_error.lower()


def test_pipeline_embedding_failure_cleans_partial_chunks(test_db):
    doc_id = uuid.uuid4()
    doc = Document(
        id=doc_id,
        owner_id=uuid.uuid4(),
        filename="sample.txt",
        file_type="txt",
        file_size=150,
        content_hash="dummy_hash_embed_fail",
        r2_object_key=f"documents/{doc_id}/sample.txt",
        status="processing",
    )
    test_db.add(doc)
    test_db.commit()

    content = b"Paragraph 1 containing important context.\n\nParagraph 2 with more information for chunking."

    with patch("backend.app.services.ingestion.SessionLocal", return_value=test_db), \
         patch("backend.app.services.gemini.GeminiService.embed_documents", side_effect=RuntimeError("Gemini API network timeout")):
        
        process_document_pipeline(
            document_id=doc_id,
            file_bytes=content,
            filename="sample.txt",
            file_type="txt",
        )

        doc = test_db.scalar(select(Document).where(Document.id == doc_id))
        assert doc.status == "failed"
        assert doc.processing_error is not None

        # Verify no partial chunks remain in database
        chunks = test_db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == doc_id)).all()
        assert len(chunks) == 0


def test_pipeline_reingestion_replaces_old_chunks(test_db):
    doc_id = uuid.uuid4()
    doc = Document(
        id=doc_id,
        owner_id=uuid.uuid4(),
        filename="repeat.txt",
        file_type="txt",
        file_size=200,
        content_hash="dummy_hash_repeat",
        r2_object_key=f"documents/{doc_id}/repeat.txt",
        status="ready",
    )
    test_db.add(doc)
    test_db.commit()

    # Pre-populate with an obsolete chunk
    old_chunk = DocumentChunk(
        document_id=doc_id,
        chunk_index=99,
        content="Obsolete content from old version",
        embedding=[0.0] * 768,
        source_location="Old source",
    )
    test_db.add(old_chunk)
    test_db.commit()

    content = b"This is the updated first paragraph with detailed system architecture.\n\nThis is the updated second paragraph containing configuration instructions."

    with patch("backend.app.services.ingestion.SessionLocal", return_value=test_db), \
         patch("backend.app.services.gemini.GeminiService.embed_documents", return_value=[[0.5] * 768] * 2):
        
        process_document_pipeline(
            document_id=doc_id,
            file_bytes=content,
            filename="repeat.txt",
            file_type="txt",
        )

        doc = test_db.scalar(select(Document).where(Document.id == doc_id))
        assert doc.status == "ready"

        # Old chunk should be deleted and replaced with new ones
        chunks = test_db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == doc_id)).all()
        assert len(chunks) > 0
        assert all(c.chunk_index != 99 for c in chunks)
        assert all("Obsolete" not in c.content for c in chunks)


def test_pipeline_dimension_mismatch_fails_and_cleans_up(test_db):
    doc_id = uuid.uuid4()
    doc = Document(
        id=doc_id,
        owner_id=uuid.uuid4(),
        filename="bad_dim.txt",
        file_type="txt",
        file_size=150,
        content_hash="dummy_hash_dim",
        r2_object_key=f"documents/{doc_id}/bad_dim.txt",
        status="processing",
    )
    test_db.add(doc)
    test_db.commit()

    content = b"Valid technical documentation that has enough characters to pass validation checks."

    # Return vectors with dimension 512 instead of configured 768
    with patch("backend.app.services.ingestion.SessionLocal", return_value=test_db), \
         patch("backend.app.services.gemini.GeminiService.embed_documents", return_value=[[0.1] * 512]):
        
        process_document_pipeline(
            document_id=doc_id,
            file_bytes=content,
            filename="bad_dim.txt",
            file_type="txt",
        )

        doc = test_db.scalar(select(Document).where(Document.id == doc_id))
        assert doc.status == "failed"
        assert "dimension" in doc.processing_error.lower()

        # No chunks saved
        chunks = test_db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == doc_id)).all()
        assert len(chunks) == 0


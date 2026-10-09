"""
Unit tests for LangChain document loaders, adapters, metadata preservation,
and temporary file lifecycle safety.
Adheres to PRD Section 6.4, 7.1, 8.
"""

import os
import io
import docx
import pypdf
import pytest
from langchain_core.documents import Document

from backend.app.rag.loaders import (
    get_langchain_loader,
    LineAwareTextLoader,
    StructuredDocxLoader,
    HeadingAwareMarkdownLoader,
    temporary_file_from_bytes,
)
from backend.app.rag.extractors import (
    get_extractor,
    langchain_doc_to_text_unit,
    langchain_docs_to_text_units,
    TextUnit,
)
from backend.app.core.errors import AppException


def test_temporary_file_cleanup_on_success():
    """Verify temporary files are deleted after successful block execution."""
    recorded_path = None
    with temporary_file_from_bytes(b"sample content", suffix=".txt") as temp_path:
        recorded_path = temp_path
        assert os.path.exists(temp_path)
        with open(temp_path, "r", encoding="utf-8") as f:
            assert f.read() == "sample content"

    assert not os.path.exists(recorded_path)


def test_temporary_file_cleanup_on_exception():
    """Verify temporary files are deleted even if an exception is raised."""
    recorded_path = None
    with pytest.raises(RuntimeError):
        with temporary_file_from_bytes(b"content to fail", suffix=".txt") as temp_path:
            recorded_path = temp_path
            assert os.path.exists(temp_path)
            raise RuntimeError("Forced error during processing")

    assert not os.path.exists(recorded_path)


def test_langchain_doc_to_text_unit_adapter():
    """Verify LangChain Document transforms accurately to TextUnit."""
    # 1. PDF 0-indexed page to 1-based page_number
    pdf_doc = Document(
        page_content="PDF page content",
        metadata={"source": "test.pdf", "page": 0, "total_pages": 3},
    )
    unit = langchain_doc_to_text_unit(pdf_doc)
    assert unit is not None
    assert unit.content == "PDF page content"
    assert unit.page_number == 1
    assert unit.source_location == "Page 1"

    # 2. Section heading metadata
    sec_doc = Document(
        page_content="Section content",
        metadata={"source": "guide.md", "section_title": "Getting Started", "source_location": "Section: Getting Started"},
    )
    sec_unit = langchain_doc_to_text_unit(sec_doc)
    assert sec_unit is not None
    assert sec_unit.section_title == "Getting Started"
    assert sec_unit.source_location == "Section: Getting Started"

    # 3. Line range metadata
    line_doc = Document(
        page_content="Line content",
        metadata={"source": "notes.txt", "section_title": "Document Body", "source_location": "Lines 10–25"},
    )
    line_unit = langchain_doc_to_text_unit(line_doc)
    assert line_unit is not None
    assert line_unit.source_location == "Lines 10–25"


def test_pypdf_loader_multi_page():
    """Verify PyPDFLoader parses multiple pages and retains 1-based page metadata."""
    import reportlab.pdfgen.canvas as rcanvas
    buf = io.BytesIO()
    c = rcanvas.Canvas(buf)
    c.drawString(100, 750, "LangChain PDF Page 1 Content")
    c.showPage()
    c.drawString(100, 750, "LangChain PDF Page 2 Content")
    c.showPage()
    c.save()
    pdf_bytes = buf.getvalue()

    extractor = get_extractor("pdf")
    units = extractor.extract(pdf_bytes, "doc.pdf")

    assert len(units) == 2
    assert units[0].page_number == 1
    assert units[0].source_location == "Page 1"
    assert "Page 1 Content" in units[0].content
    assert units[1].page_number == 2
    assert units[1].source_location == "Page 2"
    assert "Page 2 Content" in units[1].content


def test_pypdf_loader_corrupt_file():
    """Verify corrupt PDF raises CORRUPT_DOCUMENT."""
    extractor = get_extractor("pdf")
    with pytest.raises(AppException) as exc:
        extractor.extract(b"not a valid pdf header", "corrupt.pdf")
    assert exc.value.code == "CORRUPT_DOCUMENT"
    assert exc.value.status_code == 400


def test_pypdf_loader_empty_document():
    """Verify 0-byte PDF raises EMPTY_DOCUMENT."""
    extractor = get_extractor("pdf")
    with pytest.raises(AppException) as exc:
        extractor.extract(b"", "empty.pdf")
    assert exc.value.code == "EMPTY_DOCUMENT"
    assert exc.value.status_code == 400


def test_docx_loader_headings_and_tables():
    """Verify StructuredDocxLoader preserves headings and tables in metadata and text."""
    doc = docx.Document()
    doc.add_heading("Section 1: Overview", level=1)
    doc.add_paragraph("First paragraph content under overview.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"
    table.rows[1].cells[0].text = "Val A"
    table.rows[1].cells[1].text = "Val B"

    doc.add_heading("Section 2: API Endpoints", level=1)
    doc.add_paragraph("Second paragraph under API endpoints.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = get_extractor("docx")
    units = extractor.extract(docx_bytes, "api.docx")

    assert len(units) == 2
    assert units[0].section_title == "Section 1: Overview"
    assert "First paragraph content" in units[0].content
    assert "Header 1 | Header 2" in units[0].content
    assert "Val A | Val B" in units[0].content

    assert units[1].section_title == "Section 2: API Endpoints"
    assert "Second paragraph under API endpoints." in units[1].content


def test_docx_loader_corrupt_file():
    """Verify corrupt DOCX raises CORRUPT_DOCUMENT."""
    extractor = get_extractor("docx")
    with pytest.raises(AppException) as exc:
        extractor.extract(b"corrupt non-zip docx bytes", "corrupt.docx")
    assert exc.value.code == "CORRUPT_DOCUMENT"
    assert exc.value.status_code == 400


def test_markdown_loader_with_code_blocks_and_headings():
    """Verify HeadingAwareMarkdownLoader preserves code blocks without false heading splits."""
    md_content = """# System Setup
Initialize the configuration using the CLI.

```bash
# This is a comment inside a code block, NOT a heading
docker run -d postgres
```

## Security Best Practices
Always set secure secrets in production.
"""
    extractor = get_extractor("md")
    units = extractor.extract(md_content.encode("utf-8"), "setup.md")

    assert len(units) == 2
    assert units[0].section_title == "System Setup"
    assert "# This is a comment inside a code block, NOT a heading" in units[0].content
    assert units[1].section_title == "Security Best Practices"
    assert "secure secrets" in units[1].content


def test_markdown_loader_empty():
    """Verify empty Markdown raises EMPTY_DOCUMENT."""
    extractor = get_extractor("md")
    with pytest.raises(AppException) as exc:
        extractor.extract(b"   \n\n  ", "empty.md")
    assert exc.value.code == "EMPTY_DOCUMENT"


def test_text_loader_line_ranges():
    """Verify LineAwareTextLoader preserves 1-based line ranges for each paragraph."""
    txt_content = "Line 1 paragraph 1\nLine 2 paragraph 1\n\nLine 4 paragraph 2\nLine 5 paragraph 2"
    extractor = get_extractor("txt")
    units = extractor.extract(txt_content.encode("utf-8"), "sample.txt")

    assert len(units) == 2
    assert units[0].source_location == "Lines 1–2"
    assert "Line 1 paragraph 1" in units[0].content
    assert units[1].source_location == "Lines 4–5"
    assert "Line 4 paragraph 2" in units[1].content


def test_text_loader_single_line():
    """Verify single line returns 'Line X' location."""
    txt_content = "Single line document."
    extractor = get_extractor("txt")
    units = extractor.extract(txt_content.encode("utf-8"), "single.txt")

    assert len(units) == 1
    assert units[0].source_location == "Line 1"


def test_text_loader_empty():
    """Verify empty text file raises EMPTY_DOCUMENT."""
    extractor = get_extractor("txt")
    with pytest.raises(AppException) as exc:
        extractor.extract(b"    ", "empty.txt")
    assert exc.value.code == "EMPTY_DOCUMENT"


def test_unsupported_file_type():
    """Verify unsupported file type raises UNSUPPORTED_FILE_TYPE."""
    with pytest.raises(AppException) as exc:
        get_extractor("xlsx")
    assert exc.value.code == "UNSUPPORTED_FILE_TYPE"
    assert exc.value.status_code == 415


def test_get_langchain_loader_factory():
    """Verify get_langchain_loader returns appropriate BaseLoader instances."""
    with temporary_file_from_bytes(b"sample", suffix=".txt") as p_txt:
        l_txt = get_langchain_loader("txt", p_txt)
        assert isinstance(l_txt, LineAwareTextLoader)

    with temporary_file_from_bytes(b"# Heading", suffix=".md") as p_md:
        l_md = get_langchain_loader("md", p_md)
        assert isinstance(l_md, HeadingAwareMarkdownLoader)

    with pytest.raises(AppException) as exc:
        get_langchain_loader("unknown", "some/path")
    assert exc.value.code == "UNSUPPORTED_FILE_TYPE"
